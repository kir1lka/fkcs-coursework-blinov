"""Complete the retained coursework through conclusion and references.

Does not run Multisim. Input is the 28-page checkpoint, not the output.
"""
from pathlib import Path
import ast, argparse, copy, math, json, bisect
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
ASSETS=ROOT/'deliverables/figures_final'
FONT=Path('C:/Windows/Fonts')
ASSETS.mkdir(exist_ok=True)
tree=ast.parse((ROOT/'scripts/extend_sections_22_23.py').read_text(encoding='utf-8-sig'))
exec(compile(ast.Module(body=[n for n in tree.body if isinstance(n,(ast.FunctionDef,ast.ClassDef))],type_ignores=[]),'helpers','exec'))
ap=argparse.ArgumentParser();ap.add_argument('--source',required=True);args=ap.parse_args()
doc=Document(args.source)
assert not any(p.text.startswith('3. Результаты') for p in doc.paragraphs)
base=copy.deepcopy(doc.paragraphs[124]._p.pPr);capbase=copy.deepcopy(doc.paragraphs[126]._p.pPr)
SEQ=[13,2,14,11,0,12,9,5,7,8,4,3,10,6]
H=math.log(2)*95300*1e-5;T=math.log(2)*(4320+2*95300)*1e-5
FIRST=math.log(3)*(4320+95300)*1e-5
edges=[FIRST+k*T for k in range(13)];STOP=edges[-1]

def nextq(q):
    a,b,c,d=[bool(q&(1<<k)) for k in [3,2,1,0]]
    bits=[(not a and not b) or (not a and c) or (a and b and not d),
          (not a and not c and d) or (not a and c and not d) or (not b and not c) or (not b and not d),
          (not a and not b and c) or (not a and b and not c) or (a and b and d) or (a and c and not d),
          (not a and b and not c) or (not b and not c and d) or (b and not d)]
    return sum(int(v)<<(3-k) for k,v in enumerate(bits))
assert [nextq(q) for q in SEQ]==SEQ[1:]+SEQ[:1]
assert nextq(1)==13 and nextq(15)==2

def box(c,xy,title,lines):
    c.box(xy);x1,y1,x2,y2=xy;c.text((x1+x2)/2,y1+45,title,34,bold=True)
    for k,s in enumerate(lines):c.text((x1+x2)/2,y1+108+54*k,s,28)
def wire(c,pts,label=None,pos=None):
    c.line(pts,fill='#b51f24',width=4)
    if label:c.text(*pos,label,29,fill='#b51f24')
c=Canvas(1800,1210)
box(c,(80,270,410,540),'Logic_D',['Q / NOT_Q →','D3…D0'])
box(c,(650,270,1030,540),'D_Counter',['4 D-триггера','Q3…Q0; NOT_Q','P_RES → код D'])
box(c,(650,55,1030,180),'HEX_DISPLAY',['Q3…Q0'])
box(c,(1250,300,1690,630),'DC1 · 4/14',['Q / NOT_Q → Y','14 выходов','Probe на каждом Y'])
box(c,(80,810,410,1085),'P_RES_GEN',['RC + Шмитт','74LS07; S_RST','рис. 2.19'])
box(c,(650,810,1030,1105),'CLK_GEN',['LM555 + 74LS04','RUN + STOP_UNIT','START; рис. 2.20'])
box(c,(1250,850,1690,1105),'XLA1',['1–14 ← выходы Y','C ← CLK','15, 16 свободны'])
wire(c,[(410,355),(650,355)],'D3…D0',(530,315));c.text(520,415,'4 Probe',27)
wire(c,[(840,270),(840,180)])
wire(c,[(1030,405),(1250,405)],'Q / NOT_Q',(1140,365));c.dot(1120,405)
wire(c,[(1120,405),(1120,220),(245,220),(245,270)],'Обратная связь',(460,190))
wire(c,[(1120,405),(1120,945),(1030,945)]);c.text(1140,775,'Q → STOP',27,anchor='lm')
wire(c,[(840,810),(840,540)],'CLK',(780,670))
wire(c,[(410,970),(650,970)],'P_RES',(530,930))
wire(c,[(1470,630),(1470,850)],'Y0,Y2…Y14',(1500,715))
c.text(1165,1020,'CLK',28);wire(c,[(1200,1020),(1250,1020)])
c.text(900,1160,'Общая земля; питание +5 В. Одноимённые метки — одна электрическая цепь.',29)
c.save('fig_3_1_system.png')

# Exact event times, including the first charging interval and held final clock.
c=Canvas(1900,1080);x0=165;scale=76;lo=-1.;hi=20.
def xx(t):return x0+(t-lo)*scale
c.text(950,30,'Расчётная диаграмма запуска, пересчёта и остановки',35,bold=True)
for t in range(0,21,2):
    c.line([(xx(t),140),(xx(t),905)],fill='#d7d7d7',width=1);c.text(xx(t),940,str(t),28)
intervals=[lo]+edges+[hi]
for j,q in enumerate(SEQ):c.text((xx(intervals[j])+xx(intervals[j+1]))/2,100,f'{q:X}',30,bold=True)
c.text(82,100,'HEX',28)
def trace(i,label,events):
    y=165+i*88;c.text(90,y+22,label,29);pts=[]
    for j,(t,v) in enumerate(events):
        end=events[j+1][0] if j+1<len(events) else hi
        pts.extend([(xx(t),y+48-43*v),(xx(end),y+48-43*v)])
    c.line(pts,fill='#174a90' if label in ['RUN','STOP'] else '#b51f24',width=4)
trace(0,'START',[(lo,0),(0,1),(.2,0)])
trace(1,'RUN',[(lo,0),(0,1),(STOP,0)])
clock_events=[(lo,1),(0,0)]
for k,t in enumerate(edges):
    clock_events.append((t,1))
    if k<12:clock_events.append((t+H,0))
trace(2,'CLK',clock_events)
for k,b in enumerate([3,2,1,0]):trace(3+k,f'Q{b}',[(lo,(13>>b)&1)]+[(t,(SEQ[j+1]>>b)&1) for j,t in enumerate(edges)])
trace(7,'STOP',[(lo,0),(STOP,1)])
c.line([(xx(STOP),135),(xx(STOP),905)],fill='#207940',width=3)
c.text(950,995,f'От START: первый фронт {FIRST:.4f} с; остановка {STOP:.4f} с; далее HEX = 6.'.replace('.',','),30)
c.text(950,1045,'t, с. Сброс снят до START. Задержки логики не показаны; CLK после остановки = 1.',29)
c.save('fig_3_2_timing.png')

# Replace the provisional source paragraph with a reference to the final bibliography.
for p in doc.paragraphs:
    if p.text.startswith('Технические источники:'):
        p.clear();runfmt(p.add_run('Расчёт генератора и выбор элементов выполнены по технической документации [1–5]; средства наблюдения в Multisim описаны в руководстве [6].'),12)
        p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT

heading('3. Результаты синтеза управляющего устройства',1,True)
heading('3.1 Схемная реализация управляющего устройства',2)
para('На основе синтезированных блоков разработана функциональная схема управляющего устройства (рисунок 3.1). ПСУ формирует заданную последовательность четырёхразрядных кодов, дешифратор преобразует их в четырнадцать управляющих сигналов, а генератор задаёт моменты переключения и прекращает счёт в состоянии 6.')
picture('fig_3_1_system.png','Рисунок 3.1 – Расчётная функциональная схема управляющего устройства',15.3)
para('На схеме используются ранее полученные блоки Logic_D, D_Counter, DC1, CLK_GEN и P_RES_GEN. Линия обратной связи передаёт прямые и инверсные выходы регистра на комбинационную логику. Четыре Probe показывают значения D3…D0, четырнадцать Probe — активный управляющий выход; HEX-индикатор отображает текущее состояние.')
para('XLA1 записывает выходы Y по карте рисунка 2.11; для общей диаграммы он перенастраивается, как описано в 3.2. Старый тестовый генератор отключается: выходы двух генераторов нельзя соединять между собой.')

para('Связи между функциональными блоками приведены в таблице 3.1. Они дополняют поэлементные схемы раздела 2 и определяют объединение блоков в общую модель.').paragraph_format.page_break_before=True
para('Таблица 3.1 – Межблочные соединения управляющего устройства',after=5)
table(['Цепь','Источник','Приёмники'],[
['D3…D0','Logic_D','Входы D регистра; 4 Probe'],['Q3…Q0','D_Counter','Logic_D, DC1, HEX; выбранные разряды STOP_UNIT'],['NOT_Q3…NOT_Q0','D_Counter','Logic_D, DC1; выбранные разряды STOP_UNIT'],['CLK','74LS04 после LM555','4 триггера ПСУ; U_CTRL:B; XLA1:C'],['P_RES','P_RES_GEN','Асинхронная установка ПСУ; /CLR U_CTRL:B; AND2 сброса RUN'],['START','Цифровой источник запуска','CLK U_CTRL:A'],['RUN','U_CTRL:A:5','LM555:4 /RESET; контроль XSC'],['NOT_STOP','U_CTRL:B:8','AND2 совместно с P_RES'],['Y0,Y2,…,Y14','DC1','14 Probe и 14 каналов XLA1']], [3.3,4.8,8.4])
para('Регистр ПСУ содержит четыре D-триггера в двух корпусах 74AS74. Ещё один корпус используется для RUN и STOP. Дешифратор реализуется восемью AND3 и шестью AND4, логика остановки — одним AND4 и одним AND2. Времязадающая часть содержит LM555, инвертор 74LS04, R1 = 4,32 кОм, R2 = 95,3 кОм и C1 = 10 мкФ.',before=8)
para('Начальная установка формируется RC-цепью R3 = 10 кОм, C4 = 10 мкФ и двумя инверторами Шмитта 74HC14. Буфер 74LS07 с подтяжкой R4 = 1 кОм обеспечивает нагрузочную способность линии P_RES [3–5]. Питание всех блоков составляет 5 В, земля общая; у каждого корпуса предусматривается развязывающий конденсатор 100 нФ.')
para('Работа ПСУ подтверждена моделями Logic_D и D_Counter. Для общей схемы остаётся проверить запуск, дешифрацию и остановку по STOP.')

heading('3.2 Временные диаграммы сигналов устройства управления',2,True)
para('Рисунок 3.2 построен по функциям ПСУ и номиналам генератора. Отсчёт начинается с START ↑ после снятия P_RES. До запуска: код D (1101), RUN = STOP = 0, CLK = 1.')
picture('fig_3_2_timing.png','Рисунок 3.2 – Расчётная временная диаграмма сигналов управляющего устройства',16)
para('При START ↑ устанавливается RUN = 1, а CLK сначала падает. Первый положительный фронт возникает приблизительно через 1,0944 с при исходно разряженном C1; далее фронты следуют через 1,351082 с. Состояния меняются в порядке D–2–E–B–0–C–9–5–7–8–4–3–A–6.')
para('На тринадцатом фронте происходит переход A → 6 и устанавливается STOP. Время от START равно 1,0944 + 12 · 1,351082 ≈ 17,3074 с. RUN сбрасывается; CLK = 1, HEX = 6 и Y6 = 1 сохраняются. Задержки логики на расчётной диаграмме не показаны.')
para('Для этой записи XLA1 использует внутреннюю выборку 100 Гц: каналы 1–4 → Q3…Q0, 5 → START, 6 → RUN, 7 → STOP, 8 → CLK; C отключён. Запись продолжается и после остановки CLK. Диаграмма Y снимается отдельно по рисунку 2.11 [6].')

heading('3.3 Работа схемы',2,True)
para('После включения питания P_RES_GEN формирует низкий уровень начальной установки. На разрядах Q3, Q2 и Q0 активируется /PRE, на Q1 — /CLR: ПСУ принимает код 1101. Оба триггера управления сбрасываются. После заряда C4 сигнал P_RES становится высоким, и схема ожидает START.')
para('Фронт START устанавливает RUN. Таймер начинает колебания, инвертор формирует CLK. До каждого положительного фронта Logic_D вычисляет следующий код; фронт одновременно записывает D3…D0 в четыре триггера, обеспечивая синхронный пересчёт [1, 2].')
para('В каждом установившемся рабочем состоянии DC1 активирует один выход Yi, индекс которого совпадает с десятичным входным кодом. Для 1101 включён Y13, для 0010 — Y2, для конечного 0110 — Y6. HEX-индикатор показывает соответственно D, 2 и 6.')
para('В состоянии 10 активен E10. На следующем фронте ПСУ принимает 6, а триггер STOP записывает единицу. NOT_STOP через AND2 сбрасывает RUN. В этот момент OUT таймера уже низкий, поэтому его сброс сохраняет высокий CLK без дополнительного положительного фронта. Состояние 6 удерживается до общей начальной установки.')
para('Для нового прохода S_RST замыкается не менее чем на 5 с: ПСУ возвращается в 13, STOP и RUN сбрасываются, C1 разряжается. После размыкания S_RST и снятия P_RES подаётся новый фронт START. Запуск без сброса невозможен: активный STOP удерживает RUN в нуле.')
para('Проверка минимизированной логики на исключённых кодах дала переходы 1 → 13 и 15 → 2. При разрешённом тактировании возврат в рабочую последовательность занимает один такт, но до него дешифратор может выдавать несколько единиц. Поэтому штатный запуск всегда выполняется после P_RES.')

heading('Заключение',1,True)
para('В курсовой работе выполнен логический синтез управляющего устройства для последовательности 13–2–14–11–0–12–9–5–7–8–4–3–10–6. Построены таблица переходов и карты Карно, получены минимизированные функции входов четырёх D-триггеров 74AS74. Правильность рабочего цикла ПСУ подтверждена ранее выполненным моделированием и приведённой диаграммой логического анализатора.')
para('Синтезирован дешифратор 4/14 с активным единичным уровнем выходов. Проверка всех четырнадцати рабочих кодов показала единственность активного управляющего сигнала. Для исключённых кодов 1 и 15 определены переходы в рабочую последовательность и отмечено отсутствие гарантированной унитарности до установления допустимого состояния.')
para('Рассчитан управляемый генератор на LM555 с выходным инвертором. При номиналах R1 = 4,32 кОм, R2 = 95,3 кОм и C1 = 10 мкФ расчётная длительность высокого уровня равна 0,660569 с, период — 1,351082 с. Отклонение, обусловленное округлением номиналов, составляет менее 0,09 %; допуски компонентов и погрешности модели учитываются отдельно.')
para('Разработаны цепи начальной установки, запуска и остановки на состоянии 6. Регистрация признака предшествующего кода обеспечивает остановку после тринадцатого переключения от начального состояния 13. Подготовлены общая функциональная схема, межблочные соединения и расчётная диаграмма работы устройства.')
para('Полученные результаты подтверждают корректность логического синтеза и расчёта временных параметров. Полная экспериментальная проверка объединённой модели, включая RC-сброс, временные параметры и отсутствие лишнего фронта при остановке, должна быть подтверждена измерениями в NI Multisim. Расчётные диаграммы не являются результатами этих измерений.')

heading('Список используемых источников литературы',1,True)
sources=[
('Texas Instruments. LM555 Timer. SNAS548D. 2015','https://www.ti.com/lit/ds/symlink/lm555.pdf'),
('Texas Instruments. Dual Positive-Edge-Triggered D-Type Flip-Flops With Clear And Preset. SN74AS74A. SDAS143C','https://www.ti.com/lit/ds/symlink/sn74as74a.pdf'),
('Texas Instruments. Hex Inverters. SN74LS04. SDLS029C','https://www.ti.com/lit/ds/symlink/sn74ls04.pdf'),
('Texas Instruments. SNx4HC14 Hex Inverters with Schmitt-Trigger Inputs. SCLS085L. 2025','https://www.ti.com/lit/ds/symlink/sn74hc14.pdf'),
('Texas Instruments. SN74LS07 Hex Buffers and Drivers With Open-Collector High-Voltage Outputs. SDLS021D. 2016','https://www.ti.com/lit/ds/symlink/sn74ls07.pdf'),
('National Instruments. Multisim User Guide. 374483A','https://download.ni.com/support/manuals/374483a.pdf')]
for i,(title,url) in enumerate(sources,1):
    p=para(f'{i}. {title}. — Текст: электронный. — URL: {url} (дата обращения: 06.09.2026).',after=12)
    p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.keep_together=True

# Add new TOC rows after existing row 2.3.3, not after the SDT (which contains body text).
for sdt in doc._element.xpath('//w:sdt'):
    for p in list(sdt.iter(qn('w:p'))):
        if ''.join(t.text or '' for t in p.iter(qn('w:t'))).startswith('2.3.3 '):
            prev=p
            for title,pg,level in [('3. Результаты синтеза управляющего устройства',29,0),('3.1 Схемная реализация управляющего устройства',29,1),('3.2 Временные диаграммы сигналов устройства управления',31,1),('3.3 Работа схемы',32,1),('Заключение',33,0),('Список используемых источников литературы',34,0)]:
                q=doc.add_paragraph();q.paragraph_format.left_indent=Cm(level*.35);q.paragraph_format.right_indent=Cm(.5);q.paragraph_format.space_after=Pt(5);q.paragraph_format.line_spacing=1.15
                q.paragraph_format.tab_stops.add_tab_stop(Cm(17.1),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
                runfmt(q.add_run(title+'\t'+str(pg)),12);prev.addnext(q._p);prev=q._p
            break
doc.save(ROOT/'deliverables/Курсовая_Блинов_текущая.docx')
report={'sequence':SEQ,'transitions':{q:nextq(q) for q in range(16)},'first_edge_seconds':FIRST,'period_seconds':T,'stop_seconds':STOP,'edges_to_stop':len(edges),'multisim_full_system_verified':False}
(ASSETS/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
