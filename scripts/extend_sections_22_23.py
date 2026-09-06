"""Append verified calculations and original diagrams to the retained coursework.

Uses only bundled python-docx and Pillow. Never automates Multisim.
The --source argument must be a pre-extension DOCX; output defaults to canonical.
"""
from pathlib import Path
import argparse, copy, math, json, re, hashlib, io
from zipfile import ZipFile, ZIP_DEFLATED
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / 'deliverables/figures_22_23'
ASSETS.mkdir(parents=True, exist_ok=True)
SEQ = [13,2,14,11,0,12,9,5,7,8,4,3,10,6]
OUTPUTS = sorted(SEQ)
DC = [1,15]
GRAY = [0,1,3,2]
P = {}
for n in OUTPUTS:
    bits = [(3-i,(n >> (3-i)) & 1) for i in range(4)]
    neighbors = [x for x in DC if (n ^ x).bit_count()==1]
    P[n] = [(b,v) for b,v in bits if not neighbors or (1<<b)!=(n^neighbors[0])]
def eval_y(n,x): return int(all(((x>>b)&1)==v for b,v in P[n]))
assert all([n for n in OUTPUTS if eval_y(n,x)] == [x] for x in SEQ)
assert sum(len(v)==3 for v in P.values())==8
TIMING = dict(C=10e-6, R1=4320, R2=95300)
TIMING['high'] = math.log(2)*TIMING['R2']*TIMING['C']
TIMING['low'] = math.log(2)*(TIMING['R1']+TIMING['R2'])*TIMING['C']
TIMING['period'] = TIMING['high']+TIMING['low']
TIMING['frequency'] = 1/TIMING['period']

FONT = Path('C:/Windows/Fonts')
def font(n=27,bold=False): return ImageFont.truetype(str(FONT/('timesbd.ttf' if bold else 'times.ttf')), n)
class Canvas:
    def __init__(self,w,h):
        self.im=Image.new('RGB',(w,h),'white'); self.d=ImageDraw.Draw(self.im)
    def text(self,x,y,s,size=27,fill='black',anchor='mm',bold=False): self.d.text((x,y),str(s),font=font(size,bold),fill=fill,anchor=anchor)
    def line(self,pts,fill='black',width=3): self.d.line(pts,fill=fill,width=width)
    def box(self,xy,fill='white',outline='black',width=3): self.d.rectangle(xy,fill=fill,outline=outline,width=width)
    def dot(self,x,y): self.d.ellipse((x-5,y-5,x+5,y+5),fill='black')
    def save(self,name): self.im.save(ASSETS/name); return ASSETS/name

# Truth table, including explicitly unspecified input codes.
c=Canvas(1800,690); widths=[105,60,60,60,60]+[99]*14; x=[0]
for w in widths: x.append(x[-1]+w)
heads=['Код','Q3','Q2','Q1','Q0']+[f'Y{v}' for v in OUTPUTS]
for r,n in enumerate([None]+SEQ+DC):
    yy=20+r*38
    vals=heads if n is None else [str(n)]+list(f'{n:04b}')+[('X' if n in DC else str(int(v==n))) for v in OUTPUTS]
    for j,v in enumerate(vals):
        fill=('#eeeeee' if r==0 else ('#e4efdf' if j>=5 and v=='1' else 'white'))
        c.box((x[j]+12,yy,x[j+1]+12,yy+38),fill=fill,width=1)
        c.text((x[j]+x[j+1])/2+12,yy+19,v,25,bold=r==0)
c.save('fig_2_8_decoder_table.png')

# Two shared Karnaugh maps: every label denotes the sole required one of Y_n.
c=Canvas(1680,650)
colors=['#d62828','#2374ab','#6a4c93','#328044']
for k,dc in enumerate(DC):
    ox=95+k*830; oy=125; cw=160; ch=105
    c.text(ox+320,30,f'{"а" if k==0 else "б"}) Склейки с набором {dc:04b}',31,bold=True)
    c.text(ox-40,oy-37,'Q3Q2',24); c.text(ox+280,oy-68,'Q1Q0',25)
    for j,g in enumerate(GRAY): c.text(ox+(j+.5)*cw,oy-27,f'{g:02b}',27)
    for i,g in enumerate(GRAY):
        c.text(ox-32,oy+(i+.5)*ch,f'{g:02b}',27)
        for j,h in enumerate(GRAY):
            n=4*g+h
            c.box((ox+j*cw,oy+i*ch,ox+(j+1)*cw,oy+(i+1)*ch),width=2)
            c.text(ox+(j+.5)*cw,oy+(i+.5)*ch,'X' if n in DC else f'Y{n}',28)
    adj=[n for n in OUTPUTS if (n^dc).bit_count()==1]
    for z,n in enumerate(adj):
        a=(GRAY.index(n//4),GRAY.index(n%4)); b=(GRAY.index(dc//4),GRAY.index(dc%4)); col=colors[z]
        # Insets distinguish the overlapping groups without changing cell membership.
        pad=9+z*7
        if abs(a[0]-b[0])+abs(a[1]-b[1])==1:
            xx0=ox+min(a[1],b[1])*cw+pad; xx1=ox+(max(a[1],b[1])+1)*cw-pad
            yy0=oy+min(a[0],b[0])*ch+pad; yy1=oy+(max(a[0],b[0])+1)*ch-pad
            c.d.rounded_rectangle((xx0,yy0,xx1,yy1),radius=24,outline=col,width=4)
        else:
            for rr,cc in [a,b]: c.d.rounded_rectangle((ox+cc*cw+pad,oy+rr*ch+pad,ox+(cc+1)*cw-pad,oy+(rr+1)*ch-pad),radius=20,outline=col,width=4)
    c.text(ox+320,590,'Через границу: контуры одного цвета — одна группа',23)
c.save('fig_2_9_karnaugh.png')

def net_name(b,v): return f'Q{b}' if v else f'NOT_Q{b}'
c=Canvas(1600,1300)
c.text(800,32,'DC1  ·  восемь AND3 и шесть AND4',35,bold=True)
for i,n in enumerate(OUTPUTS):
    col=i//7; row=i%7; ox=40+800*col; yy=100+row*169
    c.text(ox+395,yy-24,f'U{i+1}  AND{len(P[n])}',25,fill='#174a90')
    c.box((ox+335,yy+2,ox+460,yy+121),outline='#174a90')
    c.text(ox+397,yy+60,'&',50,fill='#174a90')
    for j,(b,v) in enumerate(P[n]):
        y=yy+19+j*27
        c.text(ox+55,y,net_name(b,v),28,anchor='lm',fill='#b51f24')
        c.line([(ox+215,y),(ox+335,y)],fill='#b51f24')
    c.line([(ox+460,yy+60),(ox+595,yy+60)],fill='#b51f24')
    c.text(ox+655,yy+60,f'Y{n}',31,fill='#b51f24',bold=True)
c.text(800,1272,'Одноимённые метки — одна цепь. NOT_Qn берётся с инверсного выхода триггера.',27)
c.save('fig_2_10_decoder_scheme.png')

c=Canvas(1600,620)
c.box((50,170,350,400)); c.text(200,220,'ПСУ',36,bold=True);c.text(200,280,'Logic_D',30);c.text(200,335,'D_Counter',30)
c.box((570,170,830,400));c.text(700,245,'DC1',36,bold=True);c.text(700,310,'4/14',32)
c.line([(350,250),(570,250)],width=5);c.text(460,200,'Q / NOT_Q',26)
c.box((1080,70,1540,260));c.text(1310,100,'XLA1',33,bold=True);c.text(1310,145,'1–4: Q3, Q2, Q1, Q0',28);c.text(1310,188,'5–11: Y0,Y2,Y3,Y4,Y5,Y6,Y7',25);c.text(1310,230,'C: CLK',26)
c.box((1080,330,1540,520));c.text(1310,360,'XLA2',33,bold=True);c.text(1310,405,'1–4: Q3, Q2, Q1, Q0',28);c.text(1310,448,'5–11: Y8,Y9,Y10,Y11,Y12,Y13,Y14',24);c.text(1310,490,'C: CLK',26)
c.line([(830,230),(960,230),(960,160),(1080,160)],width=4)
c.line([(830,340),(960,340),(960,420),(1080,420)],width=4)
c.text(800,580,'Оба анализатора: External (C). Общая земля и тот же CLK, что у ПСУ.',28)
c.save('fig_2_11_decoder_test.png')

c=Canvas(1740,1040); xx=160; step=105; y0=100
c.text(880,28,'Расчётная диаграмма установившихся логических уровней',33,bold=True)
for k,n in enumerate(SEQ):
    c.text(xx+(k+.5)*step,70,f'{n:X}',28,bold=True)
    c.line([(xx+k*step,90),(xx+k*step,975)],fill='#d4d4d4',width=1)
traces=[(f'Q{b}',[(n>>b)&1 for n in SEQ]) for b in [3,2,1,0]]+[(f'Y{n}',[eval_y(n,x) for x in SEQ]) for n in OUTPUTS]
for i,(label,vals) in enumerate(traces):
    y=y0+i*47;c.text(105,y+12,label,27)
    pts=[]
    for j,v in enumerate(vals): pts.extend([(xx+j*step,y+27-26*v),(xx+(j+1)*step,y+27-26*v)])
    c.line(pts,fill='#174a90' if i<4 else '#b51f24',width=3)
c.text(900,1010,'Один столбец — одно состояние ПСУ. Задержки переключения не показаны.',28)
c.save('fig_2_12_decoder_timing.png')

c=Canvas(1300,560);c.box((445,65,835,505));c.text(640,260,'LM555',49,bold=True)
c.d.arc((590,15,690,115),0,180,fill='black',width=4)
left=['1  GND','2  TRIG','3  OUT','4  /RESET'];right=['8  VCC','7  DISCH','6  THRESH','5  CTRL']
for j in range(4):
    y=145+95*j;c.line([(360,y),(445,y)]);c.line([(835,y),(920,y)])
    c.text(320,y,left[j],33,anchor='rm');c.text(950,y,right[j],33,anchor='lm')
c.text(640,35,'Вид сверху · DIP-8',30)
c.save('fig_2_13_lm555_pinout.png')

# Schematic uses named nodes, so crossings cannot create ambiguous shorts.
c=Canvas(1660,1050)
c.text(830,30,'GEN_555  ·  питание +5 В  ·  общий GND',35,bold=True)
c.box((690,150,1070,780));c.text(880,385,'LM555CN',39,bold=True)
for y,label in [(230,'7  DISCH'),(420,'6  THRESH'),(490,'2  TRIG'),(670,'5  CTRL')]: c.text(720,y,label,30,anchor='lm')
c.text(880,178,'8  VCC',29);c.line([(880,150),(880,90)]);c.text(880,73,'+5 V',29)
c.text(880,746,'1  GND',29);c.line([(880,780),(880,850)]);c.text(880,885,'GND',29)
c.text(1038,310,'3  OUT',29,anchor='rm');c.line([(1070,310),(1270,310)])
c.d.polygon([(1270,267),(1270,353),(1350,310)],outline='#174a90',width=4);c.d.ellipse((1350,300,1370,320),outline='#174a90',width=3)
c.line([(1370,310),(1590,310)]);c.text(1470,270,'CLK',31,bold=True);c.text(1320,218,'74LS04',29);c.text(1275,377,'1 → 2',25)
c.text(1038,575,'4  /RESET',29,anchor='rm');c.line([(1070,575),(1510,575)]);c.text(1300,530,'RUN (1 — работа)',30)
c.line([(320,100),(320,135)]);c.text(320,73,'+5 V',29)
c.box((297,135,343,205));c.text(155,170,'R1 = 4,32 кОм',29);c.line([(320,205),(320,270)]);c.dot(320,230);c.line([(320,230),(690,230)])
c.box((297,270,343,370));c.text(160,320,'R2 = 95,3 кОм',29)
c.line([(320,370),(320,605)]);c.dot(320,420);c.line([(320,420),(690,420)]);c.dot(570,420);c.line([(570,420),(570,490),(690,490)])
c.line([(270,605),(370,605)],width=4);c.line([(270,625),(370,625)],width=4);c.text(254,588,'+',27);c.text(140,638,'C1 = 10 мкФ',28)
c.line([(320,625),(320,850)]);c.text(320,885,'GND',29)
c.line([(690,670),(560,670),(560,710)]);c.line([(520,710),(600,710)],width=4);c.line([(520,730),(600,730)],width=4);c.line([(560,730),(560,850)]);c.text(560,885,'GND',29);c.text(545,795,'C2 10 нФ',26,anchor='rm')
c.text(1300,720,'C3 = 100 нФ: VCC–GND у LM555',28)
c.text(1300,770,'74LS04: 14 → +5 В; 7 → GND',28)
c.text(1300,820,'У TTL-корпусов: 100 нФ VCC–GND',27)
c.text(830,964,'Для отдельной проверки: RUN = +5 В. XSC: A+ → CLK, B+ → OUT, A− и B− → GND.',28)
c.save('fig_2_14_generator_scheme.png')

c=Canvas(1600,720);c.text(800,30,'Управление запуском и остановкой · непрерывный режим',34,bold=True)
c.box((450,100,880,550));c.text(665,135,'U_RUNA · 74AS74',32,bold=True)
for y,l,net in [(220,'2  D','+5 V'),(310,'3  CLK','START'),(400,'4  /PRE','+5 V'),(490,'1  /CLR','P_RES')]:
    c.text(470,y,l,30,anchor='lm');c.line([(260,y),(450,y)]);c.text(235,y,net,32,anchor='rm')
c.text(852,265,'5  Q',30,anchor='rm');c.line([(880,265),(1110,265)]);c.text(994,222,'RUN',31)
c.box((1110,150,1550,490));c.text(1330,190,'GEN_555',34,bold=True);c.text(1330,270,'4  /RESET ← RUN',30);c.text(1330,330,'3 OUT → 74LS04 → CLK',29);c.text(1330,410,'R1, R2, C1 — рис. 2.14',28)
c.text(660,590,'U_RUN: 14 → +5 В; 7 → GND; /Q (6) не подключён',28)
c.text(800,665,'P_RES = 0: RUN = 0, ПСУ = D. START: одиночный положительный фронт.',29)
c.save('fig_2_15_run_control.png')

c=Canvas(1660,650);c.text(830,28,'Установившийся режим · расчёт по номиналам R1, R2 и C1',33,bold=True)
X0=150; scale=320; T=TIMING['period']; H=TIMING['high']
for row,inv in enumerate([False,True]):
    y=160+row*210;c.text(85,y+40,'CLK' if not inv else 'OUT',31)
    pts=[]
    for n in range(3):
        for t,val in [(n*T,1),(n*T+H,1),(n*T+H,0),((n+1)*T,0)]:
            v=1-val if inv else val;pts.append((X0+t*scale,y+95-85*v))
    c.line(pts,fill='#174a90' if inv else '#b51f24',width=4)
for t,lab in [(0,'0'),(H,'0,6606'),(T,'1,3511'),(2*T,'2,7022'),(3*T,'4,0533')]:
    xx=X0+t*scale;c.line([(xx,105),(xx,495)],fill='#bbbbbb',width=1);c.text(xx,545,lab,27)
c.text(900,106,'tH ≈ 0,6606 с       tL ≈ 0,6905 с       T ≈ 1,3511 с',31)
c.text(850,610,'t, с · начало отсчёта совмещено с фронтом CLK после установления колебаний',28)
c.save('fig_2_16_generator_timing.png')

parser=argparse.ArgumentParser();parser.add_argument('--source',required=True);parser.add_argument('--output',default=str(ROOT/'deliverables/Курсовая_Блинов_текущая.docx'));args=parser.parse_args()
doc=Document(args.source)
assert not any(p.text.startswith('2.2 Синтез') for p in doc.paragraphs),'Source already extended'
base=copy.deepcopy(doc.paragraphs[124]._p.pPr)
capbase=copy.deepcopy(doc.paragraphs[126]._p.pPr)

def runfmt(r,size=14,bold=False):
    r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor(0,0,0)
    return r
def para(text='',kind='body',before=0,after=0):
    p=doc.add_paragraph();p._p.insert(0,copy.deepcopy(capbase if kind=='caption' else base))
    f=p.paragraph_format;f.space_before=Pt(before);f.space_after=Pt(after)
    if kind=='caption': f.first_line_indent=Cm(0);f.line_spacing=1.15
    runfmt(p.add_run(text));return p
def heading(text,level=2,newpage=False):
    p=para(text,before=10,after=8);p.style=doc.styles['Heading 1']
    outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),str(level-1));p._p.get_or_add_pPr().append(outline)
    p.paragraph_format.keep_with_next=True;p.paragraph_format.page_break_before=newpage;p.paragraph_format.line_spacing=1.15
    p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs: runfmt(r,14,True)
    return p
def page():
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1
    p.add_run().add_break(WD_BREAK.PAGE);p.runs[0].font.size=Pt(1)
def picture(name,caption,width=16.5):
    p=para('',kind='caption');p.paragraph_format.keep_with_next=True
    p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(4)
    p.add_run().add_picture(str(ASSETS/name),width=Cm(width))
    dr=p._p.xpath('.//wp:docPr')[0];dr.set('descr',caption)
    q=para(caption,'caption',after=8);q.paragraph_format.keep_together=True
    return q
def mr(s):
    r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=s;r.append(t);return r
def symbol(name,sub=None,neg=False):
    el=mr(name)
    if sub is not None:
        subel=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(el);ss=OxmlElement('m:sub');ss.append(mr(str(sub)));subel.extend([e,ss]);el=subel
    if neg:
        bar=OxmlElement('m:bar');pr=OxmlElement('m:barPr');pos=OxmlElement('m:pos');pos.set(qn('m:val'),'top');pr.append(pos);e=OxmlElement('m:e');e.append(el);bar.extend([pr,e]);el=bar
    return el
def eq(nodes,num=None):
    p=para();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(5)
    om=OxmlElement('m:oMath')
    for el in nodes: om.append(mr(el) if isinstance(el,str) else el)
    p._p.append(om)
    if num:
        p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.tab_stops.add_tab_stop(Cm(8.2),WD_TAB_ALIGNMENT.CENTER)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(17.2),WD_TAB_ALIGNMENT.RIGHT)
        tr=OxmlElement('w:r');tr.append(OxmlElement('w:tab'));om.addprevious(tr)
        runfmt(p.add_run('\t('+num+')'))
    return p
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    for row in rows:
        cells=t.add_row().cells
        for cell,txt in zip(cells,row):cell.text=str(txt)
    if widths:
        for row in t.rows:
            for cell,w in zip(row.cells,widths):cell.width=Cm(w)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'000000');borders.append(e)
    t._tbl.tblPr.append(borders)
    hdr=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(hdr)
    for ri,row in enumerate(t.rows):
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.05
                for r in p.runs:runfmt(r,12,ri==0)
    return t

# Preserve all existing figures. Correct two captions locally without reflowing the chapter.
for p in doc.paragraphs:
    if p.text=='Рисунок 2.3 – Функциональная схема блока ':
        runfmt(p.add_run('Logic_D'))
    if p.text.startswith('Расчетная временная диаграмма сигнала синхронизации CLK'):
        for r in p.runs:
            if 'Расчетная' in r.text:r.text=r.text.replace('Расчетная','Полученная в Multisim',1);break
    if p.text.startswith('Рисунок 2.7 – Расчетная временная'):
        for r in p.runs:r.text=r.text.replace('Расчетная временная','Временная')

heading('2.2 Синтез дешифратора выходных состояний ПСУ',2,True)
heading('2.2.1 Условия задачи',3)
para('1. Тип дешифратора состояний ПСУ — DC 4/14 (4/16).')
para('2. Выходной код дешифратора — унитарный, активный уровень — единица.')
para('3. Индикация выходных состояний — светодиодная, с помощью Probe.')
heading('2.2.2 Порядок синтеза',3)
para('Дешифратор преобразует четырёхразрядный код Q3…Q0 в один из четырнадцати управляющих сигналов. Индекс выхода Yi соответствует десятичному значению состояния ПСУ. Например, при коде 1101 активен Y13, а при коде 0010 — Y2. Выходы Y1 и Y15 не используются. Таблица соответствия приведена на рисунке 2.8.')
picture('fig_2_8_decoder_table.png','Рисунок 2.8 – Таблица входных и выходных состояний дешифратора',17)
para('В рабочих состояниях единица присутствует ровно на одном выходе. Коды 0001 и 1111 исключены из цикла и при минимизации приняты как факультативные наборы X. Это позволяет сократить число входов части элементов И, но не гарантирует унитарный код на двух неиспользуемых наборах.')

page()
para('Минимизация выполнена по картам Карно с порядком кодов Грея 00, 01, 11, 10. Каждая метка Yi обозначает единственный обязательный единичный набор соответствующей функции; остальные рабочие клетки для этой функции равны нулю. Склейки с X показаны на рисунке 2.9.')
picture('fig_2_9_karnaugh.png','Рисунок 2.9 – Минимизация функций дешифратора: а) с набором 0001; б) с набором 1111',16.7)
para('Полученные минимальные дизъюнктивные нормальные формы имеют вид:')
for i,n in enumerate(OUTPUTS):
    factors=[]
    for j,(b,v) in enumerate(P[n]):
        if j:factors.append(' · ')
        factors.append(symbol('Q',b,not v))
    eq([symbol('Y',n),' = ']+factors,f'2.{i+1}')

page()
para('Функциональная схема DC1 содержит восемь трёхвходовых и шесть четырёхвходовых элементов И. Прямые и инверсные сигналы берутся с регистра D_Counter. В схеме рисунка 2.10 одинаковые названия цепей означают электрическое соединение; выход каждого элемента является отдельным управляющим сигналом.')
picture('fig_2_10_decoder_scheme.png','Рисунок 2.10 – Расчётная функциональная схема дешифратора DC1',16.5)
para('В NI Multisim трёхвходовые функции реализуются элементами AND3, четырёхвходовые — AND4. Для реализации на TTL допустимы три корпуса 74LS11 и три корпуса 74LS21; один элемент 74LS11 останется свободным. Неиспользуемые входы необходимо подключить к определённому логическому уровню.')

page()
heading('2.2.3 Временные диаграммы входных и выходных сигналов дешифратора',3)
para('Для проверки DC1 к его входам подключаются Q3…Q0 и NOT_Q3…NOT_Q0 работающего ПСУ. Поскольку четыре входных и четырнадцать выходных сигналов не помещаются одновременно на шестнадцати каналах одного логического анализатора, используются два прибора XLA1 и XLA2. Карта подключения приведена на рисунке 2.11.')
picture('fig_2_11_decoder_test.png','Рисунок 2.11 – Расчётная схема подключения приборов для проверки DC1')
para('Вход C обоих приборов соединяется с тактовой цепью CLK, а источник синхронизации выбирается External (C). Для обзорной диаграммы допускается использовать уже проверенный тестовый генератор ПСУ. Частота тестирования не изменяет таблицу истинности дешифратора.')
para('Для каждого установившегося состояния проверяется соответствие активного выхода коду на HEX-индикаторе. Выходы должны включаться в порядке Y13, Y2, Y14, Y11, Y0, Y12, Y9, Y5, Y7, Y8, Y4, Y3, Y10, Y6. Затем последовательность повторяется.')
para('Отдельная проверка всех шестнадцати кодов выполняется четырьмя цифровыми переключателями с инверторами или генератором слов. При коде 1 активируются Y0, Y3, Y5 и Y9; при коде 15 — Y7, Y11, Y13 и Y14. Это следствие выбранной минимизации, а не ошибка таблицы рабочего цикла.')

page()
para('На рисунке 2.12 приведена расчётная диаграмма одного полного цикла. Каждый столбец соответствует одному установившемуся состоянию; длительность столбца определяется периодом тактового сигнала.')
picture('fig_2_12_decoder_timing.png','Рисунок 2.12 – Расчётные временные диаграммы сигналов Q3…Q0 и выходов дешифратора',17)
para('Перебор всех четырнадцати рабочих кодов подтверждает единственность активного выхода. Диаграмма построена по логическим функциям (2.1)–(2.14); для подтверждения работы схемной модели требуется сопоставить её со снимками логических анализаторов NI Multisim.')
para('При переключении нескольких разрядов ПСУ задержки реальных элементов могут создавать кратковременные выбросы на выходах комбинационного дешифратора. На расчётной диаграмме они не показаны. Поэтому проверка унитарности относится к установившимся уровням, а управляющие воздействия следует оценивать после завершения переходных процессов.')

heading('2.3 Синтез управляющего генератора',2,True)
heading('2.3.1 Условия задачи',3)
para('1. Управляемый генератор — на базе микросхемы LM555.')
para('2. Длительность высокого уровня выходного импульса tH = 0,66 с.')
para('3. Период повторения импульсов T = 1,35 с.')
heading('2.3.2 Расчёт времязадающей цепи генератора импульсов',3)
para('Для формирования тактовой последовательности используется LM555 в автоколебательном режиме. Питание таймера и TTL-логики принято равным 5 В. Расположение выводов LM555 показано на рисунке 2.13 в соответствии с документацией Texas Instruments «LM555 Timer», раздел 5.')
picture('fig_2_13_lm555_pinout.png','Рисунок 2.13 – Расположение выводов микросхемы LM555',12.5)
para('Выводы 2 и 6 соединяются с времязадающим конденсатором. При достижении напряжения 2/3 питания включается разряд через вывод 7; при снижении до 1/3 питания начинается заряд. На выводе 3 формируется прямоугольный сигнал. Низкий уровень на входе /RESET (вывод 4) останавливает генерацию и устанавливает низкий уровень OUT.')

page()
para('Длительность низкого уровня на требуемом выходе CLK и коэффициент заполнения определяются по заданным значениям:')
eq([symbol('t','L'),' = T − ',symbol('t','H'),' = 1,35 − 0,66 = 0,69 с'], '2.15')
eq(['d = ',symbol('t','H'),' / T = 0,66 / 1,35 ≈ 0,4889 = 48,89 %'], '2.16')
para('В стандартной астабильной схеме LM555 длительность высокого уровня OUT больше длительности низкого. Поэтому для получения заполнения менее 50 % на выходе устанавливается инвертор 74LS04. После инвертирования высокий уровень CLK соответствует разряду C1, а низкий — заряду. Схема с рассчитанными номиналами показана на рисунке 2.14.')
picture('fig_2_14_generator_scheme.png','Рисунок 2.14 – Расчётная схема генератора на LM555 с инвертором',17)
para('На выводе OUT таймера требуются высокий уровень 0,69 с и низкий уровень 0,66 с. После 74LS04 эти интервалы меняются местами. Точку измерения CLK следует выбирать именно после инвертора.')

page()
para('Для стандартной схемы LM555 используются зависимости времени заряда и разряда из документации производителя, раздел 7.4.2. При расчёте коэффициент 0,693 уточнён до ln 2. Для инвертированного выхода получаем:')
eq([symbol('t','H'),' = ln 2 · ',symbol('R','2'),' · ',symbol('C','1')], '2.17')
eq([symbol('t','L'),' = ln 2 · (',symbol('R','1'),' + ',symbol('R','2'),') · ',symbol('C','1')], '2.18')
eq(['T = ln 2 · (',symbol('R','1'),' + 2',symbol('R','2'),') · ',symbol('C','1')], '2.19')
eq(['f = 1 / T = 1 / 1,35 ≈ 0,74074 Гц'], '2.20')
para('Принимаем C1 = 10 мкФ. Требуемые сопротивления составляют:')
eq([symbol('R','2'),' = 0,66 / (ln 2 · 10⁻⁵) ≈ 95 217,87 Ом'], '2.21')
eq([symbol('R','1'),' = (0,69 − 0,66) / (ln 2 · 10⁻⁵) ≈ 4 328,09 Ом'], '2.22')
para('Выбираем ближайшие номиналы ряда E96: R1 = 4,32 кОм и R2 = 95,3 кОм. При идеальных номиналах и порогах таймера получаем значения таблицы 2.1. Погрешность рассчитана относительно технического задания.')
para('Таблица 2.1 – Расчётные временные параметры генератора',after=5)
table(['Параметр','Задание','По номиналам','Отклонение'],[
    ['Высокий уровень CLK','0,6600 с',f'{TIMING["high"]:.6f} с'.replace('.',','),f'{100*(TIMING["high"]/.66-1):+.4f} %'.replace('.',',')],
    ['Низкий уровень CLK','0,6900 с',f'{TIMING["low"]:.6f} с'.replace('.',','),f'{100*(TIMING["low"]/.69-1):+.4f} %'.replace('.',',')],
    ['Период CLK','1,3500 с',f'{TIMING["period"]:.6f} с'.replace('.',','),f'{100*(TIMING["period"]/1.35-1):+.4f} %'.replace('.',',')],
    ['Частота CLK','0,740741 Гц',f'{TIMING["frequency"]:.6f} Гц'.replace('.',','),f'{100*(TIMING["frequency"]/(1/1.35)-1):+.4f} %'.replace('.',',')]], [5.6,3.3,4.1,3.8])
para('Отклонение из-за округления номиналов менее 0,09 %. Оно не является полной погрешностью устройства: фактические интервалы также зависят от допусков R и C, утечки конденсатора, порогов компараторов и напряжения насыщения разрядного транзистора.',before=8)
para('Например, при резисторах ±1 % и конденсаторе ±10 % только произведение RC меняется от 0,891 до 1,111 номинала. Дополнительная погрешность LM555 в этот интервал не включена. Для точной настройки сначала подбирают R2 по длительности высокого уровня CLK, затем R1 по периоду.')

page()
para('Для переноса схемы в NI Multisim применяется карта соединений таблицы 2.2. Все обозначения относятся к рисунку 2.14. Времязадающий конденсатор C1 подключается положительным выводом к общей точке TRIG/THRESH, отрицательным — к земле.')
para('Таблица 2.2 – Карта соединений генератора',after=5)
table(['Элемент или вывод','Соединение'],[
    ['LM555: 8 VCC; 1 GND','8 → +5 В; 1 → общая земля'],
    ['LM555: 7 DISCH','К общей точке R1 и R2'],
    ['LM555: 2 TRIG и 6 THRESH','Объединить; к нижнему выводу R2 и плюсу C1'],
    ['R1 = 4,32 кОм','Между +5 В и выводом 7'],
    ['R2 = 95,3 кОм','Между выводом 7 и узлом выводов 2, 6'],
    ['C1 = 10 мкФ','Плюс к узлу 2, 6; минус к GND'],
    ['C2 = 10 нФ','Между выводом 5 CTRL и GND'],
    ['C3 = 100 нФ','Между 8 VCC и 1 GND, рядом с таймером'],
    ['LM555: 4 /RESET','К RUN; при отдельной проверке — к +5 В'],
    ['LM555: 3 OUT','На вход 1 первого инвертора 74LS04'],
    ['74LS04: 2','Выход CLK на тактовые входы ПСУ'],
    ['74LS04: 14; 7','14 → +5 В; 7 → GND; между ними 100 нФ'],
    ['Остальные входы 74LS04','Подключить к GND; соответствующие выходы оставить свободными'],
    ['Осциллограф XSC1','A+ → CLK; B+ → OUT; A− и B− → GND']], [6.2,10.6])
para('Управление генератором выполняется сигналом RUN. При RUN = 1 таймер формирует импульсы; при RUN = 0 вывод OUT удерживается в нуле. Поскольку за таймером установлен инвертор, остановленному генератору соответствует CLK = 1, а не CLK = 0. Это необходимо учитывать при начальной установке и остановке ПСУ.',before=8)

page()
para('Разрешение работы хранится в дополнительном D-триггере 74AS74: D = 1, /PRE = 1, CLK = START, /CLR = P_RES. Его прямой выход RUN подключён к выводу 4 LM555. Схема управления приведена на рисунке 2.15.')
picture('fig_2_15_run_control.png','Рисунок 2.15 – Расчётная схема управления генератором на D-триггере',15.5)
para('При P_RES = 0 триггер разрешения сброшен, генератор остановлен, ПСУ установлен в код D (1101). После снятия P_RES положительный фронт START записывает единицу в триггер разрешения. Таймер запускается; CLK сначала падает, поэтому D-триггеры ПСУ не переключаются.')
para('При первом запуске C1 заряжается почти от нуля. Задержка первого фронта составляет ln 3 · (R1 + R2) · C1 ≈ 1,0944 с и отличается от установившейся паузы. Заданный период измеряется по последующим импульсам.')
para('Повторный P_RES останавливает генератор и возвращает ПСУ в начальный код. Сброс только таймера при низком CLK может вызвать лишний пересчёт: после инвертора возникнет положительный фронт. Поэтому остановка совмещена с общей начальной установкой. Автоматическая остановка по конечному состоянию не вводится; замкнутый цикл сохранён.')
para('START и P_RES задаются цифровыми источниками. Для физических кнопок необходима защита от дребезга. Входы свободной секции 74AS74 фиксируются; одновременный ноль на /PRE и /CLR запрещён.')

page()
heading('2.3.3 Временные диаграммы выходных сигналов генератора',3)
para('Расчётная диаграмма сигналов OUT и CLK приведена на рисунке 2.16. Начало отсчёта выбрано по положительному фронту CLK после установления колебаний. На выходе инвертора высокий уровень составляет около 0,6606 с, низкий — 0,6905 с, а период — 1,3511 с.')
picture('fig_2_16_generator_timing.png','Рисунок 2.16 – Расчётные временные диаграммы OUT и CLK',16)
para('В NI Multisim канал A осциллографа XSC1 подключается к CLK, канал B — к OUT таймера. Для обзора можно начать с масштаба 0,5 с/деление. После установления колебаний моделирование приостанавливается.')
para('Для измерения импульса первый курсор устанавливается на положительный фронт CLK, второй — на следующий отрицательный: ожидается около 0,6606 с. Для измерения периода курсоры ставятся на соседние положительные фронты CLK: ожидается около 1,3511 с. На OUT длительности высокого и низкого уровней меняются местами.')
para('Расчёт соответствует заданным параметрам с отклонением номиналов менее 0,09 %. Схемная модель требует проверки в Multisim: расчётная диаграмма не является записью осциллографа. Для отчёта сохраняются схема генератора и два снимка с курсорами — измерение импульса и периода.')
p=para('Технические источники: Texas Instruments, LM555 Timer (SNAS548D), разделы 5 и 7.4.2, https://www.ti.com/lit/ds/symlink/lm555.pdf; Hex Inverters (SDLS029C), https://www.ti.com/lit/ds/symlink/sn74ls04.pdf; D-Type Flip-Flops With Clear And Preset (SDAS143C), https://www.ti.com/lit/ds/symlink/sn74as74a.pdf. Дата обращения: 06.09.2026.',before=8)
p.paragraph_format.line_spacing=1.0
for r in p.runs:runfmt(r,12)

# Replace only the existing TOC content control with a static source-styled TOC.
# Page numbers are finalized against the Word export by the QA command.
toc=doc._element.xpath('//w:sdt[w:sdtPr/w:docPartObj]')
if not toc:
    toc=[s for s in doc._element.xpath('//w:sdt') if any('TOC' in str(x) for x in s.xpath('.//w:instrText/text()'))]
toc_entries=[
('1 Техническое задание на проектирование',4,0),('1.1 Цель и содержание курсовой работы',4,1),('1.2 Вариант задания для моделирования разработанного устройства',4,1),
('2 Синтез управляющего устройства',5,0),('2.1 Синтез пересчётного устройства',5,1),('2.1.1 Порядок синтеза',5,2),('2.1.3 Временные диаграммы сигналов ПСУ',12,2),
('2.2 Синтез дешифратора выходных состояний ПСУ',13,1),('2.2.1 Условия задачи',13,2),('2.2.2 Порядок синтеза',13,2),('2.2.3 Временные диаграммы сигналов дешифратора',16,2),
('2.3 Синтез управляющего генератора',18,1),('2.3.1 Условия задачи',18,2),('2.3.2 Расчёт времязадающей цепи генератора импульсов',18,2),('2.3.3 Временные диаграммы выходных сигналов генератора',23,2)]
if toc:
    content=toc[0].find(qn('w:sdtContent'))
    if content is not None:
        # Source SDT extends past the TOC into the technical-assignment chapter.
        # Only its first six paragraphs are TOC entries; retain the rest intact.
        tail=[copy.deepcopy(e) for e in list(content)[6:]]
        for child in list(content): content.remove(child)
        for txt,pg,level in toc_entries:
            p=doc.add_paragraph();p.paragraph_format.left_indent=Cm(level*.35);p.paragraph_format.right_indent=Cm(.5);p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1.15
            p.paragraph_format.tab_stops.add_tab_stop(Cm(17.1),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
            runfmt(p.add_run(txt+'\t'+str(pg)),12)
            content.append(p._p)
        for e in tail:content.append(e)

doc.save(args.output)
# Preserve all opaque original package parts byte-for-byte. python-docx otherwise
# normalizes XML declarations in untouched styles, footers, settings and metadata.
with ZipFile(args.source) as original, ZipFile(args.output) as authored:
    parts={n:authored.read(n) for n in authored.namelist()}
    for n in original.namelist():
        if n not in {'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}:
            parts[n]=original.read(n)
buf=io.BytesIO()
with ZipFile(buf,'w',ZIP_DEFLATED) as z:
    for n,data in parts.items():z.writestr(n,data)
Path(args.output).write_bytes(buf.getvalue())
report={'sequence':SEQ,'decoder_terms':P,'unused_outputs':{x:[n for n in OUTPUTS if eval_y(n,x)] for x in DC},'timing':TIMING,'source_sha256':hashlib.sha256(Path(args.source).read_bytes()).hexdigest(),'assets':[p.name for p in ASSETS.glob('*.png')]}
(ASSETS/'calculation_check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
